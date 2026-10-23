"""엑셀(.xlsx)·CSV 한 행 = 협상안 요청 하나 (#236).

첫 행은 머리글이다. 머리글은 OfferProposalRequest 필드 이름이나 한글 별칭(협력사, RFQ, 목표단가 등)을 쓴다.
행마다 따로 검증해서 잘못된 행은 오류로 보고하고 나머지 행은 계속 처리한다.
"""

import csv
import io
import zipfile
from collections.abc import Awaitable, Callable
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, ValidationError

from .request import OfferProposalRequest

# 한글 머리글 → 필드. 공백·괄호는 무시하고 비교한다
ALIASES = {
    '협력사': 'supplier', '협력사id': 'supplier', '공급사': 'supplier', 'rfq': 'rfq_id', 'rfqid': 'rfq_id',
    '기준일': 'as_of', '작성일': 'as_of', '기준': 'intent', '목표단가': 'target_price', '최대수용가': 'max_price',
    '지급일수': 'payment_days', '지급기한': 'payment_days', '연동': 'price_linkage', '연동약정': 'price_linkage',
    '원재료비중': 'raw_material_ratio', '에너지비중': 'energy_ratio', '계약기간': 'contract_days',
    '계약기간일': 'contract_days', '미연동합의': 'linkage_waived_by_agreement', '비고': 'notes', '납품비고': 'notes',
}
FIELDS = [f for f in OfferProposalRequest.model_fields if f not in ('supplier_size', 'buyer_size', 'sender', 'transaction')]
TRUE = {'예', 'y', 'yes', 'o', 'true', '1', '포함', '있음', '연동'}
FALSE = {'아니오', '아니요', 'n', 'no', 'x', 'false', '0', '없음', '미포함'}
RATIO_FIELDS = ('raw_material_ratio', 'energy_ratio')
TEMPLATE_HEADERS = ['협력사', 'RFQ', '기준일', '목표단가', '최대수용가', '지급일수', '연동', '원재료비중', '에너지비중',
                    '계약기간', '미연동합의', '비고']


class RowResult(BaseModel):
    row: int                                   # 엑셀 행 번호 (머리글이 1행)
    request: OfferProposalRequest | None = None
    error: str | None = None


def header_field(name: Any) -> str | None:
    key = ''.join(ch for ch in str(name or '').strip().lower() if ch not in ' _-()·')
    if key in ALIASES:
        return ALIASES[key]
    plain = str(name or '').strip()
    return plain if plain in FIELDS else None


def coerce(field: str, value: Any) -> Any:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, str):
        value = value.strip()
    if field in ('price_linkage', 'linkage_waived_by_agreement'):
        if isinstance(value, bool):
            return value
        text = str(value).lower()
        if text in TRUE:
            return True
        if text in FALSE:
            return False
        raise ValueError(f'예/아니오로 적어 주세요: {value}')
    if field in RATIO_FIELDS:
        text = str(value).replace('%', '').strip()
        number = float(text)
        return number / 100 if (str(value).endswith('%') or number > 1) else number
    if field in ('target_price', 'max_price', 'payment_days', 'contract_days'):
        number = float(str(value).replace(',', '').replace('원', '').replace('일', ''))
        if number != int(number):
            raise ValueError(f'정수로 적어 주세요: {value}')
        return int(number)
    if field == 'as_of':
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        return str(value)
    return str(value) if field in ('supplier', 'rfq_id', 'notes', 'intent') else value


def read_table(data: bytes, filename: str) -> list[list[Any]]:
    if filename.lower().endswith('.csv'):
        text = data.decode('utf-8-sig')
        return [row for row in csv.reader(io.StringIO(text))]
    sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def parse_rows(data: bytes, filename: str, default_as_of: str) -> list[RowResult]:
    table = read_table(data, filename)
    if not table:
        return []
    fields = [header_field(h) for h in table[0]]
    unknown = [str(h) for h, f in zip(table[0], fields) if h not in (None, '') and f is None]
    out = []
    for number, row in enumerate(table[1:], start=2):
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in row):
            continue
        if unknown:
            out.append(RowResult(row=number, error=f'알 수 없는 머리글: {", ".join(unknown)}'))
            continue
        try:
            body = {f: coerce(f, v) for f, v in zip(fields, row) if f}
            body = {k: v for k, v in body.items() if v is not None}
            body.setdefault('as_of', default_as_of)
            out.append(RowResult(row=number, request=OfferProposalRequest.model_validate(body, strict=False)))
        except (ValueError, ValidationError) as exc:
            out.append(RowResult(row=number, error=short_error(exc)))
    return out


def short_error(exc: Exception) -> str:
    if isinstance(exc, ValidationError):
        return '; '.join(f'{".".join(map(str, e["loc"])) or "행"}: {e["msg"]}' for e in exc.errors())
    return str(exc)


def template_bytes() -> bytes:
    """머리글과 예시 2행이 있는 양식."""
    wb = Workbook()
    ws = wb.active
    ws.title = '협상안 요청'
    ws.append(TEMPLATE_HEADERS)
    ws.append(['SUP-0055', 'RFQ-2026-004', '2026-10-23', None, None, 90, '아니오', None, None, 365, '아니오',
               '월 2회 분할 납품'])
    ws.append(['SUP-0060', 'RFQ-2026-002', '2026-10-23', 1750, 2050, 60, None, '42%', None, 365, None, None])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


SUMMARY_COLUMNS = ['row', 'supplier', 'rfq_id', 'status', 'unit_price', 'payment_days', 'linkage', 'warnings',
                   'input_needed', 'content_hash', 'files', 'error']


def summary_csv(rows: list[dict]) -> bytes:
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=SUMMARY_COLUMNS, extrasaction='ignore')
    writer.writeheader()
    writer.writerows(rows)
    return buffer.getvalue().encode('utf-8-sig')


Producer = Callable[[OfferProposalRequest], Awaitable[tuple[dict, dict[str, bytes]]]]


async def run_batch(rows: list[RowResult], produce: Producer) -> tuple[list[dict], dict[str, bytes]]:
    """행마다 produce(요청) → (요약 한 줄, {파일 이름: 바이트}). 한 행이 실패해도 다음 행을 처리한다."""
    summary, files = [], {}
    for r in rows:
        base = {'row': r.row, 'supplier': r.request.supplier if r.request else '',
                'rfq_id': r.request.rfq_id if r.request else ''}
        if r.request is None:
            summary.append(base | {'status': 'invalid', 'error': r.error})
            continue
        try:
            line, produced = await produce(r.request)
        except Exception as exc:  # noqa: BLE001  한 행의 도구 오류가 일괄 처리 전체를 멈추지 않게 한다
            summary.append(base | {'status': 'error', 'error': str(exc)[:300]})
            continue
        prefixed = {f'{r.row:03d}_{name}': data for name, data in produced.items()}
        files |= prefixed
        summary.append(base | line | {'files': ' '.join(prefixed)})
    return summary, files


def zip_bytes(files: dict[str, bytes], summary: list[dict]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data)
        z.writestr('summary.csv', summary_csv(summary))
    return buffer.getvalue()


def write_out_dir(out_dir: Path, files: dict[str, bytes], summary: list[dict]) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    for name, data in files.items():
        (out_dir / name).write_bytes(data)
    (out_dir / 'summary.csv').write_bytes(summary_csv(summary))
