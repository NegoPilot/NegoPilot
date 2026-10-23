"""협력사 발송용 "단가 및 거래조건 제안서" PDF·DOCX (#236).

OfferDocument의 값만 그린다. 내부 협상 값(목표 합의가, 최대 수용가, 양보 순서, 쟁점, 시장가·직전 견적 비교, 경고,
법규 검토 내부 내용)과 "AI" 표시는 넣지 않는다. PDF와 DOCX는 같은 sections()를 그린다.
"""

from io import BytesIO
from urllib.parse import quote

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

from .document import OfferDocument
from .pdf import Block, draw, won

DOCX_TYPE = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
DOCX_FONT = '맑은 고딕'


def ratio(value: float | None) -> str:
    return '-' if value is None else f'{value * 100:.0f}%'


def sections(d: OfferDocument) -> list[tuple[str, list[Block]]]:
    s = d.sender
    item = [['RFQ', d.rfq_id], ['품목', d.part_name]]
    if d.vehicle:
        item.append(['차종', d.vehicle])
    item.append(['연간 물량', f'{d.annual_volume:,}개'])
    linkage = ('포함' if d.price_linkage else '미연동 합의' if d.linkage_waived_by_agreement
               else '협의 후 결정' if d.linkage_needs_input else '해당 없음')
    return [
        ('단가 및 거래조건 제안서', [('table', [
            ['문서 번호', f'{d.document_no} (v{d.version})'], ['작성일', d.issued_on.isoformat()],
            ['수신', f'{d.supplier_name} ({d.supplier_id}) 귀중'],
            ['발신', f'{s.company} {s.department} {s.contact_name} (전화 {s.phone}, {s.email})'],
            ['제목', d.subject]])]),
        ('인사말', [('p', d.cover_letter)]),
        ('대상 품목', [('table', item)]),
        ('제안 조건', [('table', [
            ['단가', won(d.unit_price)], ['계약 기간', f'{d.contract_days}일'],
            ['예상 계약금액', f'{won(d.contract_amount)} (연간 물량 기준, 부가가치세 별도)'],
            ['지급 조건', d.payment_terms()], ['납품대금 연동', linkage]])]),
        ('납품대금 연동', [('list', d.linkage_lines())]),
        ('납품 및 비고', [('p', d.delivery_notes or '-')]),
        ('회신 요청', [('p', f'{d.response_due.isoformat()}까지 회신 부탁드립니다. '
                         f'문의: {s.contact_name} ({s.phone}, {s.email})')]),
        ('서명', [('table', [['발신', f'{s.company} {s.department}'], ['담당자', f'{s.contact_name} (서명/인)'],
                            ['수신 확인', f'{d.supplier_name} (서명/인)']])]),
    ]


def render_pdf(d: OfferDocument) -> bytes:
    return draw(sections(d), d.subject)


def render_docx(d: OfferDocument) -> bytes:
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = DOCX_FONT
    style.font.size = Pt(10)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), DOCX_FONT)
    for n, (name, blocks) in enumerate(sections(d)):
        heading = doc.add_heading(name, level=0 if n == 0 else 2)
        if n == 0:
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for kind, value in blocks:
            if kind == 'p':
                doc.add_paragraph(value)
            elif kind == 'list':
                for line in value:
                    doc.add_paragraph(line, style='List Bullet')
            else:
                table = doc.add_table(rows=0, cols=len(value[0]))
                table.style = 'Table Grid'
                for row in value:
                    cells = table.add_row().cells
                    for cell, content in zip(cells, row):
                        cell.text = str(content)
    doc.core_properties.title = d.subject
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def docx_text(data: bytes) -> str:
    doc = Document(BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    parts += [c.text for t in doc.tables for row in t.rows for c in row.cells]
    return '\n'.join(parts)


def content_disposition(filename: str) -> str:
    ascii_name = filename.encode('ascii', 'replace').decode().replace('?', '_')
    return f'attachment; filename="{ascii_name}"; filename*=UTF-8\'\'{quote(filename)}'
